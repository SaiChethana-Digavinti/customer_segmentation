import re
import os
import markdown
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from PIL import Image as PILImage
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable, Preformatted, Image as RLImage
)
from reportlab.pdfgen import canvas

MD_PATH = r"c:\Users\saich\Desktop\customer-segmentation\Project_documentation.md"
DOCX_PATH = r"c:\Users\saich\Desktop\customer-segmentation\Project_documentation.docx"
HTML_PATH = r"c:\Users\saich\Desktop\customer-segmentation\Project_documentation.html"
PDF_PATH = r"c:\Users\saich\Desktop\customer-segmentation\Project_documentation.pdf"

# -------------------------------------------------------------
# 1. GENERATE HTML WITH PRINT-TO-PDF STYLING
# -------------------------------------------------------------
def generate_html(md_text):
    print("Generating HTML version...")
    # Convert markdown to html with table and code extensions
    html_body = markdown.markdown(
        md_text,
        extensions=['tables', 'fenced_code', 'toc', 'sane_lists']
    )
    
    # Custom academic theme with page break styling
    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Enterprise Customer Segmentation - Technical Documentation</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
    <style>
        :root {{
            --primary: #1e3a8a;
            --secondary: #4f46e5;
            --text-main: #0f172a;
            --text-muted: #475569;
            --bg-page: #f8fafc;
            --border-color: #e2e8f0;
        }}
        * {{ box-sizing: border-box; }}
        body {{
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            background-color: var(--bg-page);
            color: var(--text-main);
            line-height: 1.7;
            margin: 0;
            padding: 0;
        }}
        .top-nav {{
            position: sticky;
            top: 0;
            background: rgba(15, 23, 42, 0.95);
            backdrop-filter: blur(10px);
            color: white;
            padding: 12px 24px;
            display: flex;
            justify-content: space-between;
            align-items: center;
            z-index: 1000;
            box-shadow: 0 4px 12px rgba(0,0,0,0.15);
        }}
        .nav-title {{
            font-size: 15px;
            font-weight: 600;
            letter-spacing: 0.5px;
        }}
        .btn-download {{
            background: linear-gradient(135deg, #4f46e5, #06b6d4);
            color: white;
            border: none;
            padding: 8px 18px;
            border-radius: 6px;
            font-weight: 600;
            font-size: 13px;
            cursor: pointer;
            transition: all 0.2s ease;
            text-decoration: none;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        }}
        .btn-download:hover {{
            transform: translateY(-1px);
            box-shadow: 0 4px 12px rgba(79, 70, 229, 0.4);
        }}
        .container {{
            max-width: 900px;
            margin: 30px auto;
            background: #ffffff;
            padding: 50px 70px;
            border-radius: 8px;
            box-shadow: 0 4px 25px rgba(0,0,0,0.06);
            border: 1px solid var(--border-color);
        }}
        h1 {{
            color: #0f172a;
            font-size: 26px;
            border-bottom: 2px solid #e2e8f0;
            padding-bottom: 8px;
            margin-top: 40px;
        }}
        h2 {{
            color: var(--primary);
            font-size: 20px;
            margin-top: 30px;
            border-bottom: 1px solid #f1f5f9;
            padding-bottom: 6px;
        }}
        h3 {{
            color: #334155;
            font-size: 16px;
            margin-top: 20px;
        }}
        p, li {{
            font-size: 14.5px;
            color: #334155;
            text-align: justify;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 25px 0;
            font-size: 13.5px;
        }}
        th, td {{
            border: 1px solid #cbd5e1;
            padding: 9px 12px;
            text-align: left;
        }}
        th {{
            background-color: #0f172a;
            color: white;
            font-weight: 600;
        }}
        tr:nth-child(even) {{
            background-color: #f8fafc;
        }}
        pre, code {{
            font-family: 'JetBrains Mono', Consolas, Monaco, monospace;
            background-color: #0f172a;
            color: #e2e8f0;
            border-radius: 6px;
        }}
        code {{
            padding: 2px 6px;
            font-size: 13px;
            background: #f1f5f9;
            color: #4338ca;
        }}
        pre {{
            padding: 16px;
            overflow-x: auto;
            font-size: 12.5px;
            line-height: 1.5;
        }}
        pre code {{
            background: transparent;
            color: inherit;
            padding: 0;
        }}
        blockquote {{
            border-left: 4px solid var(--secondary);
            margin: 18px 0;
            padding: 10px 18px;
            background: #f8fafc;
            color: #475569;
            font-style: italic;
        }}
        hr {{
            border: none;
            border-top: 1px solid #e2e8f0;
            margin: 35px 0;
        }}
        img {{
            max-width: 100%;
            height: auto;
            border-radius: 8px;
            box-shadow: 0 4px 20px rgba(0,0,0,0.08);
            margin: 22px auto 8px auto;
            display: block;
            border: 1px solid #e2e8f0;
        }}
        p em {{
            display: block;
            text-align: center;
            font-size: 13px;
            color: #64748b;
            margin-top: 4px;
            margin-bottom: 20px;
        }}
        @media print {{
            .no-print, .top-nav {{ display: none !important; }}
            body {{ background: white; color: black; }}
            .container {{
                max-width: 100%;
                margin: 0;
                padding: 0;
                border: none;
                box-shadow: none;
            }}
            h1 {{
                page-break-before: always;
                break-before: page;
            }}
            h1:first-of-type {{
                page-break-before: avoid;
                break-before: avoid;
            }}
            table, pre, blockquote {{
                page-break-inside: avoid;
            }}
            @page {{
                margin: 20mm 15mm;
            }}
        }}
    </style>
</head>
<body>
    <div class="top-nav no-print">
        <div class="nav-title">📄 Enterprise Customer Segmentation & AI Marketing Platform — Project Documentation</div>
        <div>
            <button class="btn-download" onclick="window.print()">
                <svg width="15" height="15" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M6 9V2h12v7"></path><path d="M6 18H4a2 2 0 0 1-2-2v-5a2 2 0 0 1 2-2h16a2 2 0 0 1 2 2v5a2 2 0 0 1-2 2h-2"></path><rect x="6" y="14" width="12" height="8"></rect></svg>
                Save as PDF / Print Document
            </button>
        </div>
    </div>
    <div class="container">
        {html_body}
    </div>
</body>
</html>
"""
    with open(HTML_PATH, "w", encoding="utf-8") as f:
        f.write(full_html)
    print(f"HTML generated at: {HTML_PATH}")

# -------------------------------------------------------------
# 2. GENERATE MICROSOFT WORD (.DOCX) DOCUMENT
# -------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_docx(md_text):
    print("Generating DOCX version...")
    doc = Document()

    # Set page margins (1 inch)
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Style default normal font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    lines = md_text.splitlines()
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle fenced code block
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                code_text = "\n".join(code_lines)
                p = doc.add_paragraph()
                pPr = p._element.get_or_add_pPr()
                # Shading background for code block
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
                pPr.append(shd)
                run = p.add_run(code_text)
                run.font.name = 'Courier New'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(15, 23, 42)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle Markdown Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table rows
            parsed_rows = []
            for tl in table_lines:
                # Check if separator row like |:---|:---|
                if re.match(r'^\|[\s:-]+\|$', tl.replace(" ", "")):
                    continue
                cells = [c.strip() for c in tl.strip("|").split("|")]
                parsed_rows.append(cells)
            
            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = False

                for r_idx, r_data in enumerate(parsed_rows):
                    row = tbl.rows[r_idx]
                    for c_idx in range(num_cols):
                        c_text = r_data[c_idx] if c_idx < len(r_data) else ""
                        cell = row.cells[c_idx]
                        cell.text = clean_markdown_inline(c_text)
                        cell_p = cell.paragraphs[0]
                        cell_p.paragraph_format.space_after = Pt(2)
                        cell_p.paragraph_format.space_before = Pt(2)
                        if r_idx == 0:
                            set_cell_background(cell, "0F172A")
                            for run in cell_p.runs:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                                run.font.size = Pt(9.5)
                        else:
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "F8FAFC")
                            for run in cell_p.runs:
                                run.font.size = Pt(9.5)
                doc.add_paragraph() # Spacer after table
            continue

        # Handle Images ![alt](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if img_match:
            rel_path = img_match.group(2)
            full_img_path = os.path.normpath(os.path.join(r"c:\Users\saich\Desktop\customer-segmentation", rel_path))
            if os.path.exists(full_img_path):
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.paragraph_format.space_before = Pt(12)
                img_p.paragraph_format.space_after = Pt(4)
                run = img_p.add_run()
                run.add_picture(full_img_path, width=Inches(5.8))
            i += 1
            continue

        # Handle Figure Captions *Figure X.X: ...*
        if stripped.startswith("*Figure") and stripped.endswith("*"):
            cap_p = doc.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_before = Pt(2)
            cap_p.paragraph_format.space_after = Pt(14)
            run = cap_p.add_run(stripped.strip("*"))
            run.font.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            h_text = stripped[2:].strip()
            # If it's CHAPTER 2 or later, add a page break before it
            if h_text.startswith("CHAPTER") and not h_text.startswith("CHAPTER 1:"):
                doc.add_page_break()
            h = doc.add_heading(level=1)
            run = h.add_run(clean_markdown_inline(h_text))
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(8)
            i += 1
            continue

        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            h = doc.add_heading(level=2)
            run = h.add_run(clean_markdown_inline(h_text))
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            h = doc.add_heading(level=3)
            run = h.add_run(clean_markdown_inline(h_text))
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Horizontal rule
        if stripped == "---":
            # Add a subtle paragraph spacing
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Unordered list item
        if stripped.startswith("- ") or stripped.startswith("* "):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            add_formatted_text(p, item_text)
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # Numbered list item
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            item_text = m_num.group(2)
            p = doc.add_paragraph(style='List Number')
            add_formatted_text(p, item_text)
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # Blockquote
        if stripped.startswith(">"):
            quote_text = stripped.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            run = p.add_run(clean_markdown_inline(quote_text))
            run.font.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Empty line
        if not stripped:
            i += 1
            continue

        # Standard paragraph
        p = doc.add_paragraph()
        add_formatted_text(p, stripped)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        i += 1

    doc.save(DOCX_PATH)
    print(f"DOCX generated at: {DOCX_PATH}")

def clean_markdown_inline(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)
    text = text.replace('$', '')
    return text

def add_formatted_text(paragraph, text):
    """Parse bold, italic, code tokens and add runs to paragraph."""
    # Split tokens by markdown bold and code
    pattern = r'(\*\*.*?\*\*|`.*?`|\*.*?\*)'
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.font.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(79, 70, 229)
        else:
            # Normal text (clean math delimiters if any)
            clean_token = token.replace('$', '')
            paragraph.add_run(clean_token)

# -------------------------------------------------------------
# 3. GENERATE PDF DOCUMENT VIA REPORTLAB
# -------------------------------------------------------------
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        canvas.Canvas.__init__(self, *args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8.5)
        self.setFillColor(colors.HexColor("#64748b"))
        # Header (pages after page 1)
        if self._pageNumber > 1:
            self.drawString(54, 750, "Enterprise Customer Segmentation & AI Marketing Platform — Technical Documentation")
            self.setStrokeColor(colors.HexColor("#e2e8f0"))
            self.setLineWidth(0.5)
            self.line(54, 742, 558, 742)
        # Footer
        footer_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 36, footer_text)
        self.drawString(54, 36, "Confidential — Academic & Technical Project Report")
        self.setStrokeColor(colors.HexColor("#e2e8f0"))
        self.setLineWidth(0.5)
        self.line(54, 48, 558, 48)
        self.restoreState()

def generate_pdf(md_text):
    print("Generating PDF version via ReportLab...")
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=8
    )
    h1_style = ParagraphStyle(
        'DocH1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#0f172a'),
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=True
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor('#1e3a8a'),
        spaceBefore=12,
        spaceAfter=5,
        keepWithNext=True
    )
    h3_style = ParagraphStyle(
        'DocH3',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#334155'),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=5
    )
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=body_style,
        leftIndent=15,
        spaceAfter=3
    )
    code_style = ParagraphStyle(
        'DocCode',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0f172a'),
        backColor=colors.HexColor('#f1f5f9'),
        leftIndent=10,
        rightIndent=10,
        spaceBefore=6,
        spaceAfter=6
    )

    caption_style = ParagraphStyle(
        'DocCaption',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8.5,
        leading=11.5,
        alignment=1,
        textColor=colors.HexColor('#64748b'),
        spaceBefore=4,
        spaceAfter=12
    )

    story = []
    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code
        if stripped.startswith("```"):
            if in_code:
                in_code = False
                safe_code = "<br/>".join(
                    code_lines[:40]
                ).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe_code, code_style))
                story.append(Spacer(1, 4))
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            parsed_rows = []
            for tl in table_lines:
                if re.match(r'^\|[\s:-]+\|$', tl.replace(" ", "")):
                    continue
                cells = [c.strip() for c in tl.strip("|").split("|")]
                parsed_rows.append(cells)

            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                col_width = 504 / num_cols
                table_data = []
                for r_idx, r in enumerate(parsed_rows):
                    row_data = []
                    for c_idx in range(num_cols):
                        txt = r[c_idx] if c_idx < len(r) else ""
                        safe_txt = clean_markdown_inline(txt).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                        p_cell = Paragraph(f"<b>{safe_txt}</b>" if r_idx == 0 else safe_txt, body_style)
                        row_data.append(p_cell)
                    table_data.append(row_data)

                t = Table(table_data, colWidths=[col_width]*num_cols)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                    ('TOPPADDING', (0, 0), (-1, -1), 4),
                    ('LEFTPADDING', (0, 0), (-1, -1), 4),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')])
                ]))
                story.append(t)
                story.append(Spacer(1, 6))
            continue

        # Handle Images ![alt](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if img_match:
            rel_path = img_match.group(2)
            full_img_path = os.path.normpath(os.path.join(r"c:\Users\saich\Desktop\customer-segmentation", rel_path))
            if os.path.exists(full_img_path):
                try:
                    with PILImage.open(full_img_path) as pil_img:
                        orig_w, orig_h = pil_img.size
                    max_w = 480.0
                    max_h = 270.0
                    scale = min(max_w / orig_w, max_h / orig_h, 1.0)
                    w = orig_w * scale
                    h = orig_h * scale
                    story.append(Spacer(1, 8))
                    story.append(RLImage(full_img_path, width=w, height=h))
                except Exception as e:
                    print(f"Error embedding image {full_img_path} in PDF: {e}")
            i += 1
            continue

        # Handle Figure Captions *Figure X.X: ...*
        if stripped.startswith("*Figure") and stripped.endswith("*"):
            caption_text = stripped.strip("*").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(caption_text, caption_style))
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            h_text = stripped[2:].strip()
            if h_text.startswith("CHAPTER") and not h_text.startswith("CHAPTER 1:"):
                story.append(PageBreak())
            story.append(Paragraph(clean_markdown_inline(h_text), h1_style))
            i += 1
            continue

        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            story.append(Paragraph(clean_markdown_inline(h_text), h2_style))
            i += 1
            continue

        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            story.append(Paragraph(clean_markdown_inline(h_text), h3_style))
            i += 1
            continue

        if stripped == "---":
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e2e8f0"), spaceBefore=8, spaceAfter=8))
            i += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = clean_markdown_inline(stripped[2:].strip())
            safe_text = bullet_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(f"• {safe_text}", bullet_style))
            i += 1
            continue

        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            safe_text = clean_markdown_inline(m_num.group(2)).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(f"{m_num.group(1)}. {safe_text}", bullet_style))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Standard paragraph
        safe_p = clean_markdown_inline(stripped).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe_p, body_style))
        i += 1

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"PDF generated at: {PDF_PATH}")

def main():
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()

    generate_html(md_text)
    generate_docx(md_text)
    generate_pdf(md_text)
    print("\nAll 3 downloadable document formats successfully created!")

if __name__ == "__main__":
    main()
